"""LLM-facing artifact generation tools for Part 3.

These tools create and edit files/projects inside the secure workspace layer.
"""
from __future__ import annotations

from typing import Any
from pathlib import Path
import csv
from io import StringIO
from .workspace import WorkspaceManager
from .manager import ArtifactManager
from .types import ArtifactType
from .generators import generate_docx, generate_xlsx, generate_pdf, generate_pptx, extension_for

MAX_PROJECT_FILES = 100
MAX_PROJECT_PAYLOAD_BYTES = 50 * 1024 * 1024

CREATE_ARTIFACT_TOOL = {
    "type": "function",
    "function": {
        "name": "artifact_create_file",
        "description": (
            "Create one new text/code file in the current assistant workspace. "
            "Use this for plain text or source-code files such as .txt, .md, .json, .csv, .html, .css, .js, .py. "
            "DO NOT use this tool for .docx, .xlsx, .pdf, or .pptx; use the dedicated generator tool for those formats. "
            "Do not overwrite an existing file."
        ),
        "parameters": {
            "type": "object",
            "properties": {
                "filename": {"type": "string", "description": "Relative filename, e.g. index.html"},
                "content": {"type": "string", "description": "Complete UTF-8 text content of the file"},
                "artifact_type": {
                    "type": "string",
                    "enum": ["file", "document", "spreadsheet", "presentation", "image", "project"],
                    "description": "High-level artifact category; normally use file for code/text files.",
                },
            },
            "required": ["filename", "content"],
            "additionalProperties": False,
        },
    },
}

CREATE_PROJECT_TOOL = {
    "type": "function",
    "function": {
        "name": "artifact_create_project",
        "description": (
            "Create a new multi-file project in the current assistant workspace. "
            "Use this when the user requests a website or other project requiring multiple files."
        ),
        "parameters": {
            "type": "object",
            "properties": {
                "project_name": {"type": "string", "description": "Human-readable project name"},
                "files": {
                    "type": "array",
                    "items": {
                        "type": "object",
                        "properties": {
                            "path": {"type": "string"},
                            "content": {"type": "string"},
                        },
                        "required": ["path", "content"],
                        "additionalProperties": False,
                    },
                },
            },
            "required": ["project_name", "files"],
            "additionalProperties": False,
        },
    },
}

EDIT_ARTIFACT_TOOL = {
    "type": "function",
    "function": {
        "name": "artifact_edit_file",
        "description": (
            "Replace the complete UTF-8 text content of an existing artifact file. "
            "Use this only when the user asks to modify a file previously created or managed by the assistant. "
            "Do not change the artifact ID, workspace, or path."
        ),
        "parameters": {
            "type": "object",
            "properties": {
                "artifact_id": {"type": "string", "description": "Stable artifact ID of the existing file"},
                "content": {"type": "string", "description": "Complete replacement UTF-8 text content"},
            },
            "required": ["artifact_id", "content"],
            "additionalProperties": False,
        },
    },
}

LIST_ARTIFACTS_TOOL = {
    "type": "function",
    "function": {
        "name": "artifact_list",
        "description": "List existing artifacts in the current assistant workspace so the assistant can identify files to edit.",
        "parameters": {
            "type": "object",
            "properties": {},
            "additionalProperties": False,
        },
    },
}


CREATE_DOCX_TOOL = {
    "type": "function",
    "function": {
        "name": "artifact_create_document",
        "description": "Create a real Microsoft Word .docx document from a title and paragraphs. Use this tool whenever the user asks for a Word document, DOCX, or Microsoft Word file.",
        "parameters": {
            "type": "object",
            "properties": {
                "filename": {"type": "string", "description": "Output .docx filename"},
                "title": {"type": "string"},
                "paragraphs": {"type": "array", "items": {"type": "string"}},
            },
            "required": ["filename", "title", "paragraphs"],
            "additionalProperties": False,
        },
    },
}

CREATE_XLSX_TOOL = {
    "type": "function",
    "function": {
        "name": "artifact_create_spreadsheet",
        "description": "Create a real Microsoft Excel .xlsx workbook from structured sheets and rows. Use this tool whenever the user asks for an Excel spreadsheet, XLSX, or Microsoft Excel file.",
        "parameters": {
            "type": "object",
            "properties": {
                "filename": {"type": "string", "description": "Output .xlsx filename"},
                "sheets": {
                    "type": "array", "items": {
                        "type": "object",
                        "properties": {
                            "name": {"type": "string"},
                            "rows": {"type": "array", "items": {"type": "array", "items": {}}},
                        },
                        "required": ["name", "rows"],
                        "additionalProperties": False,
                    },
                },
            },
            "required": ["filename", "sheets"],
            "additionalProperties": False,
        },
    },
}

CREATE_PDF_TOOL = {
    "type": "function",
    "function": {
        "name": "artifact_create_pdf",
        "description": "Create a real PDF document from a title and paragraphs. Use this tool whenever the user asks for a PDF file.",
        "parameters": {
            "type": "object",
            "properties": {
                "filename": {"type": "string", "description": "Output .pdf filename"},
                "title": {"type": "string"},
                "paragraphs": {"type": "array", "items": {"type": "string"}},
            },
            "required": ["filename", "title", "paragraphs"],
            "additionalProperties": False,
        },
    },
}

CREATE_PPTX_TOOL = {
    "type": "function",
    "function": {
        "name": "artifact_create_presentation",
        "description": "Create a real Microsoft PowerPoint .pptx presentation. Use this tool whenever the user asks for PowerPoint, PPTX, or a presentation.",
        "parameters": {
            "type": "object",
            "properties": {
                "filename": {"type": "string", "description": "Output .pptx filename"},
                "title": {"type": "string"},
                "slides": {
                    "type": "array", "items": {
                        "type": "object",
                        "properties": {
                            "title": {"type": "string"},
                            "bullets": {"type": "array", "items": {"type": "string"}},
                        },
                        "required": ["title", "bullets"],
                        "additionalProperties": False,
                    },
                },
            },
            "required": ["filename", "title", "slides"],
            "additionalProperties": False,
        },
    },
}

CONVERT_ARTIFACT_TOOL = {
    "type": "function",
    "function": {
        "name": "artifact_convert",
        "description": (
            "Convert an existing artifact to another supported format. Use this whenever the user asks to convert, change format, export, or turn one existing artifact into another format. "
            "Create a new artifact and preserve the original. Supported conversions include "
            "DOCX to TXT/MD, TXT/MD to DOCX, CSV to XLSX, and XLSX to CSV."
        ),
        "parameters": {
            "type": "object",
            "properties": {
                "artifact_id": {"type": "string"},
                "target_format": {
                    "type": "string",
                    "enum": ["txt", "md", "docx", "csv", "xlsx"],
                },
                "filename": {
                    "type": "string",
                    "description": "Optional output filename. Extension is enforced by the conversion target.",
                },
            },
            "required": ["artifact_id", "target_format"],
            "additionalProperties": False,
        },
    },
}

ARTIFACT_TOOL_NAMES = {
    "artifact_create_file",
    "artifact_create_project",
    "artifact_edit_file",
    "artifact_list",
    "artifact_create_document",
    "artifact_create_spreadsheet",
    "artifact_create_pdf",
    "artifact_create_presentation",
    "artifact_convert",
}
ARTIFACT_TOOLS_OPENAI = [CREATE_ARTIFACT_TOOL, CREATE_PROJECT_TOOL, EDIT_ARTIFACT_TOOL, LIST_ARTIFACTS_TOOL, CREATE_DOCX_TOOL, CREATE_XLSX_TOOL, CREATE_PDF_TOOL, CREATE_PPTX_TOOL, CONVERT_ARTIFACT_TOOL]
# Claude's tool schema is compatible with Anthropic's tools format via the existing adapter shape.
ARTIFACT_TOOLS_CLAUDE = [
    {
        "name": CREATE_ARTIFACT_TOOL["function"]["name"],
        "description": CREATE_ARTIFACT_TOOL["function"]["description"],
        "input_schema": CREATE_ARTIFACT_TOOL["function"]["parameters"],
    },
    {
        "name": CREATE_PROJECT_TOOL["function"]["name"],
        "description": CREATE_PROJECT_TOOL["function"]["description"],
        "input_schema": CREATE_PROJECT_TOOL["function"]["parameters"],
    },
    {
        "name": EDIT_ARTIFACT_TOOL["function"]["name"],
        "description": EDIT_ARTIFACT_TOOL["function"]["description"],
        "input_schema": EDIT_ARTIFACT_TOOL["function"]["parameters"],
    },
    {
        "name": LIST_ARTIFACTS_TOOL["function"]["name"],
        "description": LIST_ARTIFACTS_TOOL["function"]["description"],
        "input_schema": LIST_ARTIFACTS_TOOL["function"]["parameters"],
    },
]
for _tool in (CREATE_DOCX_TOOL, CREATE_XLSX_TOOL, CREATE_PDF_TOOL, CREATE_PPTX_TOOL, CONVERT_ARTIFACT_TOOL):
    ARTIFACT_TOOLS_CLAUDE.append({
        "name": _tool["function"]["name"],
        "description": _tool["function"]["description"],
        "input_schema": _tool["function"]["parameters"],
    })


def _artifact_type(value: str | None) -> ArtifactType:
    if not value:
        return ArtifactType.FILE
    try:
        return ArtifactType(value)
    except ValueError:
        return ArtifactType.FILE


class ArtifactToolService:
    """Creates new artifacts using an isolated workspace."""

    def __init__(
        self,
        workspace_manager: WorkspaceManager,
        artifact_manager: ArtifactManager,
        workspace_id: str | None = None,
        conversation_id: str | None = None,
    ) -> None:
        self.workspace_manager = workspace_manager
        self.artifact_manager = artifact_manager
        self.workspace_id = workspace_id
        self.conversation_id = conversation_id

    def _ensure_workspace(self) -> str:
        if self.workspace_id:
            return self.workspace_id
        if self.conversation_id:
            existing = self.workspace_manager.get_by_conversation_id(self.conversation_id)
            if existing:
                self.workspace_id = existing.workspace_id
                return self.workspace_id
        workspace = self.workspace_manager.create(name="VTuber Artifacts", conversation_id=self.conversation_id)
        self.workspace_id = workspace.workspace_id
        return self.workspace_id

    def create_file(self, filename: str, content: str, artifact_type: str | None = None) -> dict[str, Any]:
        workspace_id = self._ensure_workspace()
        if not isinstance(filename, str) or not filename.strip():
            raise ValueError("filename must be a non-empty string")
        if not isinstance(content, str):
            raise ValueError("content must be a string")
        ext = Path(filename).suffix.lower()
        generator_only = {".docx": "artifact_create_document", ".xlsx": "artifact_create_spreadsheet", ".pdf": "artifact_create_pdf", ".pptx": "artifact_create_presentation"}
        if ext in generator_only:
            raise ValueError(f"Use {generator_only[ext]} for {ext} files; artifact_create_file only creates text/code files.")
        # Creation-only: the workspace layer performs an atomic create-if-absent.
        path = self.workspace_manager.create_bytes(
            workspace_id, filename, content.encode("utf-8")
        )
        rel = path.relative_to(self.workspace_manager.require(workspace_id).root_path).as_posix()
        mime = _guess_mime(filename)
        artifact = self.artifact_manager.create(
            name=path.name,
            artifact_type=_artifact_type(artifact_type),
            mime_type=mime,
            workspace_id=workspace_id,
            path=rel,
            size=len(content.encode("utf-8")),
        )
        return artifact.to_dict()

    def edit_file(self, artifact_id: str, content: str) -> dict[str, Any]:
        """Replace an existing managed artifact's complete text content."""
        workspace_id = self._ensure_workspace()
        if not isinstance(artifact_id, str) or not artifact_id.strip():
            raise ValueError("artifact_id must be a non-empty string")
        if not isinstance(content, str):
            raise ValueError("content must be a string")

        artifact = self.artifact_manager.require(artifact_id)
        if artifact.workspace_id != workspace_id:
            raise PermissionError("Artifact does not belong to the current workspace")
        if not artifact.path:
            raise ValueError("Artifact has no editable file path")

        # Resolve through the secure workspace layer; never trust the stored path blindly.
        old_bytes = self.workspace_manager.read_bytes(workspace_id, artifact.path)
        new_bytes = content.encode("utf-8")
        if new_bytes == old_bytes:
            # Keep metadata consistent without rewriting the file.
            self.artifact_manager.update(artifact_id, size=len(new_bytes))
            return self.artifact_manager.require(artifact_id).to_dict()

        self.workspace_manager.write_bytes(workspace_id, artifact.path, new_bytes)
        try:
            updated = self.artifact_manager.update(artifact_id, size=len(new_bytes))
        except Exception:
            # Best-effort restoration if metadata persistence fails after the file replacement.
            try:
                self.workspace_manager.write_bytes(workspace_id, artifact.path, old_bytes)
            except Exception:
                pass
            raise
        return updated.to_dict()

    def _create_generated_file(self, filename: str, data: bytes, artifact_type: ArtifactType, mime: str) -> dict[str, Any]:
        workspace_id = self._ensure_workspace()
        if len(data) > self.workspace_manager.max_file_size:
            raise ValueError("Generated file exceeds workspace file-size limit")
        path = self.workspace_manager.create_bytes(workspace_id, filename, data)
        rel = path.relative_to(self.workspace_manager.require(workspace_id).root_path).as_posix()
        try:
            artifact = self.artifact_manager.create(
                name=path.name, artifact_type=artifact_type, mime_type=mime,
                workspace_id=workspace_id, path=rel, size=len(data),
            )
        except Exception:
            try:
                self.workspace_manager.delete_file(workspace_id, rel)
            finally:
                raise
        return artifact.to_dict()

    def create_generated_file(
        self, filename: str, data: bytes, artifact_type: str, mime: str
    ) -> dict[str, Any]:
        """Persist already-generated bytes using the secure artifact workspace."""
        return self._create_generated_file(
            filename, data, _artifact_type(artifact_type), mime
        )

    def create_document(self, filename: str, title: str, paragraphs: list[str]) -> dict[str, Any]:
        expected, mime = extension_for("docx")
        if not filename.lower().endswith(expected):
            filename = f"{filename}{expected}"
        return self._create_generated_file(filename, generate_docx(title=title, paragraphs=paragraphs), ArtifactType.DOCUMENT, mime)

    def create_spreadsheet(self, filename: str, sheets: list[dict[str, Any]]) -> dict[str, Any]:
        expected, mime = extension_for("xlsx")
        if not filename.lower().endswith(expected):
            filename = f"{filename}{expected}"
        return self._create_generated_file(filename, generate_xlsx(sheets=sheets), ArtifactType.SPREADSHEET, mime)

    def create_pdf(self, filename: str, title: str, paragraphs: list[str]) -> dict[str, Any]:
        expected, mime = extension_for("pdf")
        if not filename.lower().endswith(expected):
            filename = f"{filename}{expected}"
        return self._create_generated_file(filename, generate_pdf(title=title, paragraphs=paragraphs), ArtifactType.DOCUMENT, mime)

    def create_presentation(self, filename: str, title: str, slides: list[dict[str, Any]]) -> dict[str, Any]:
        expected, mime = extension_for("pptx")
        if not filename.lower().endswith(expected):
            filename = f"{filename}{expected}"
        return self._create_generated_file(filename, generate_pptx(title=title, slides=slides), ArtifactType.PRESENTATION, mime)

    def convert_artifact(self, artifact_id: str, target_format: str, filename: str | None = None) -> dict[str, Any]:
        """Convert a managed artifact into a new artifact, preserving the original."""
        workspace_id = self._ensure_workspace()
        artifact = self.artifact_manager.require(artifact_id)
        if artifact.workspace_id != workspace_id:
            raise PermissionError("Artifact does not belong to the current workspace")
        if not artifact.path:
            raise ValueError("Artifact has no file path")
        source_path = self.workspace_manager._resolve_file(workspace_id, artifact.path)
        if source_path.is_symlink() or not source_path.exists() or not source_path.is_file():
            raise FileNotFoundError("Artifact file not found")

        source_ext = Path(artifact.name).suffix.lower()
        target = str(target_format).lower().lstrip(".")
        conversion_key = (source_ext, target)
        raw = self.workspace_manager.read_bytes(workspace_id, artifact.path)

        if conversion_key in {(".docx", "txt"), (".docx", "md")} :
            from docx import Document
            doc = Document(StringIO(raw.decode("utf-8"))) if False else Document(source_path)
            lines = []
            for paragraph in doc.paragraphs:
                text = paragraph.text
                if text:
                    lines.append(text)
            content = "\n\n".join(lines).encode("utf-8")
            default_ext = target
        elif conversion_key in {(".txt", "docx"), (".md", "docx")}:
            from docx import Document
            text = raw.decode("utf-8")
            title = Path(artifact.name).stem
            data = generate_docx(title=title, paragraphs=text.splitlines())
            content = data
            default_ext = "docx"
        elif conversion_key == (".csv", "xlsx"):
            from openpyxl import Workbook
            wb = Workbook(write_only=True)
            ws = wb.create_sheet(title="Sheet1")
            text = raw.decode("utf-8-sig")
            reader = csv.reader(StringIO(text))
            for row in reader:
                ws.append([_safe_excel_value(v) for v in row])
            out = BytesIO()
            wb.save(out)
            wb.close()
            content = out.getvalue()
            default_ext = "xlsx"
        elif conversion_key == (".xlsx", "csv"):
            from openpyxl import load_workbook
            wb = load_workbook(source_path, read_only=True, data_only=True)
            try:
                target_ws = wb.active
                out = StringIO(newline="")
                writer = csv.writer(out)
                for row in target_ws.iter_rows(values_only=True):
                    writer.writerow(["" if v is None else str(v) for v in row])
                content = out.getvalue().encode("utf-8")
            finally:
                wb.close()
            default_ext = "csv"
        else:
            raise ValueError(f"Unsupported conversion: {source_ext or 'unknown'} -> {target}")

        base = str(filename).strip() if filename else Path(artifact.name).stem + "." + default_ext
        if not base.lower().endswith("." + default_ext):
            base = Path(base).stem + "." + default_ext
        return self._create_generated_file(base, content, ArtifactType.DOCUMENT if default_ext in {"docx", "pdf"} else (ArtifactType.SPREADSHEET if default_ext == "xlsx" else ArtifactType.FILE), _guess_mime(base))

    def list_artifacts(self) -> list[dict[str, Any]]:
        workspace_id = self._ensure_workspace()
        return [artifact.to_dict() for artifact in self.artifact_manager.list(workspace_id=workspace_id)]

    def create_project(self, project_name: str, files: list[dict[str, Any]]) -> list[dict[str, Any]]:
        workspace_id = self._ensure_workspace()
        if not isinstance(project_name, str) or not project_name.strip():
            raise ValueError("project_name must be a non-empty string")
        if not isinstance(files, list) or not files:
            raise ValueError("files must be a non-empty list")
        if len(files) > MAX_PROJECT_FILES:
            raise ValueError(f"Project contains too many files (max {MAX_PROJECT_FILES})")
        # Validate the entire project before writing anything, so invalid input cannot
        # leave a half-created project.
        normalized: list[tuple[str, str]] = []
        seen: set[str] = set()
        for item in files:
            if not isinstance(item, dict):
                raise ValueError("Each project file must be an object")
            path = item.get("path")
            content = item.get("content")
            if not isinstance(path, str) or not path.strip() or not isinstance(content, str):
                raise ValueError("Each project file requires string path and content")
            # Validate path without creating it.
            candidate = self.workspace_manager._safe_relative_path(path)
            norm = candidate.as_posix()
            if norm in seen:
                raise ValueError(f"Duplicate project path: {norm}")
            seen.add(norm)
            normalized.append((norm, content))

        total_bytes = sum(len(content.encode("utf-8")) for _, content in normalized)
        if total_bytes > MAX_PROJECT_PAYLOAD_BYTES:
            raise ValueError(f"Project payload exceeds {MAX_PROJECT_PAYLOAD_BYTES} bytes")
        if total_bytes > self.workspace_manager.max_workspace_size:
            raise ValueError("Project exceeds workspace capacity")
        current_files = self.workspace_manager.file_count(workspace_id)
        if current_files + len(normalized) > self.workspace_manager.max_files_per_workspace:
            raise ValueError(
                f"Workspace file-count limit exceeded ({self.workspace_manager.max_files_per_workspace} files)"
            )

        created: list[dict[str, Any]] = []
        created_paths: list[str] = []
        created_artifact_ids: list[str] = []
        try:
            for path, content in normalized:
                # Creation-only semantics: never overwrite an existing path.
                # create_text/create_bytes use atomic O_EXCL creation in the workspace layer.
                self.workspace_manager.create_bytes(workspace_id, path, content.encode("utf-8"))
                created_paths.append(path)
                try:
                    artifact = self.artifact_manager.create(
                        name=path.rsplit("/", 1)[-1],
                        artifact_type=ArtifactType.FILE,
                        mime_type=_guess_mime(path),
                        workspace_id=workspace_id,
                        path=path,
                        size=len(content.encode("utf-8")),
                        metadata={"project_name": project_name.strip()},
                    )
                    created_artifact_ids.append(artifact.artifact_id)
                    created.append(artifact.to_dict())
                except Exception:
                    # This file was created by this operation; remove only that file.
                    try:
                        self.workspace_manager.delete_file(workspace_id, path)
                    finally:
                        raise
        except Exception:
            # Roll back only files/artifacts actually created by this operation.
            for artifact_id in reversed(created_artifact_ids):
                try:
                    self.artifact_manager.delete(artifact_id)
                except Exception:
                    pass
            for path in reversed(created_paths):
                try:
                    self.workspace_manager.delete_file(workspace_id, path)
                except Exception:
                    pass
            raise
        return created


def _guess_mime(filename: str) -> str:
    import mimetypes
    mime, _ = mimetypes.guess_type(filename)
    return mime or "application/octet-stream"
