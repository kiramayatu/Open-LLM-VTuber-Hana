"""Safe, structured generators for common document formats."""
from __future__ import annotations

from io import BytesIO
from tempfile import SpooledTemporaryFile
from typing import Any, Iterable

MAX_TEXT_BYTES = 10 * 1024 * 1024
MAX_PARAGRAPHS = 500
MAX_SHEETS = 20
MAX_ROWS_PER_SHEET = 10_000
MAX_COLS_PER_ROW = 200
MAX_SLIDES = 100
MAX_BULLETS_PER_SLIDE = 50


def _ensure_text(value: Any, field: str) -> str:
    if not isinstance(value, str):
        raise ValueError(f"{field} must be a string")
    return value


def _validate_text_budget(parts: Iterable[str]) -> None:
    total = 0
    for part in parts:
        total += len(part.encode("utf-8"))
        if total > MAX_TEXT_BYTES:
            raise ValueError(f"Generated text exceeds {MAX_TEXT_BYTES} bytes")


def generate_docx(*, title: str, paragraphs: list[str]) -> bytes:
    title = _ensure_text(title, "title").strip()
    if not title:
        raise ValueError("title cannot be empty")
    if not isinstance(paragraphs, list) or len(paragraphs) > MAX_PARAGRAPHS:
        raise ValueError(f"paragraphs must be a list with at most {MAX_PARAGRAPHS} items")
    paragraphs = [_ensure_text(p, "paragraph") for p in paragraphs]
    _validate_text_budget([title, *paragraphs])

    from docx import Document
    from docx.shared import Pt

    doc = Document()
    heading = doc.add_heading(title, level=1)
    if heading.runs:
        heading.runs[0].font.size = Pt(18)
    for paragraph in paragraphs:
        doc.add_paragraph(paragraph)
    out = BytesIO()
    doc.save(out)
    return out.getvalue()


def _safe_excel_value(value: Any) -> Any:
    # Prevent spreadsheet formula injection from LLM/user-controlled strings.
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@"):
        return "'" + value
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def generate_xlsx(*, sheets: list[dict[str, Any]]) -> bytes:
    if not isinstance(sheets, list) or not sheets or len(sheets) > MAX_SHEETS:
        raise ValueError(f"sheets must be a non-empty list with at most {MAX_SHEETS} sheets")

    from openpyxl import Workbook

    # Write-only mode keeps only the current row/cell structures in memory.
    wb = Workbook(write_only=True)
    text_bytes = 0

    try:
        for sheet_spec in sheets:
            if not isinstance(sheet_spec, dict):
                raise ValueError("Each sheet must be an object")
            name = _ensure_text(sheet_spec.get("name"), "sheet name").strip() or "Sheet"
            rows = sheet_spec.get("rows")
            if not isinstance(rows, list) or len(rows) > MAX_ROWS_PER_SHEET:
                raise ValueError(f"rows must be a list with at most {MAX_ROWS_PER_SHEET} rows")
            if len(name) > 31 or any(ch in name for ch in "[]:*?/\\"):
                raise ValueError("Invalid Excel sheet name")
            ws = wb.create_sheet(title=name)
            text_bytes += len(name.encode("utf-8"))
            if text_bytes > MAX_TEXT_BYTES:
                raise ValueError(f"Generated text exceeds {MAX_TEXT_BYTES} bytes")
            for row in rows:
                if not isinstance(row, list) or len(row) > MAX_COLS_PER_ROW:
                    raise ValueError(f"Each row must be a list with at most {MAX_COLS_PER_ROW} columns")
                safe_row = []
                for value in row:
                    safe_value = _safe_excel_value(value)
                    safe_row.append(safe_value)
                    if isinstance(safe_value, str):
                        text_bytes += len(safe_value.encode("utf-8"))
                        if text_bytes > MAX_TEXT_BYTES:
                            raise ValueError(f"Generated text exceeds {MAX_TEXT_BYTES} bytes")
                ws.append(safe_row)

        # Keep the final ZIP payload in a spooled buffer so small workbooks stay in RAM
        # while large ones spill to disk automatically.
        out = SpooledTemporaryFile(max_size=8 * 1024 * 1024, mode="w+b")
        try:
            wb.save(out)
            out.seek(0)
            data = out.read()
        finally:
            out.close()
        return data
    finally:
        try:
            wb.close()
        except Exception:
            pass

def generate_pdf(*, title: str, paragraphs: list[str]) -> bytes:
    title = _ensure_text(title, "title").strip()
    if not title:
        raise ValueError("title cannot be empty")
    if not isinstance(paragraphs, list) or len(paragraphs) > MAX_PARAGRAPHS:
        raise ValueError(f"paragraphs must be a list with at most {MAX_PARAGRAPHS} items")
    paragraphs = [_ensure_text(p, "paragraph") for p in paragraphs]
    _validate_text_budget([title, *paragraphs])

    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    from reportlab.lib.units import mm

    out = BytesIO()
    pdf = canvas.Canvas(out, pagesize=A4, pageCompression=1)
    width, height = A4
    y = height - 25 * mm
    pdf.setTitle(title[:200])
    pdf.setFont("Helvetica-Bold", 16)
    pdf.drawString(20 * mm, y, title[:120])
    y -= 12 * mm
    pdf.setFont("Helvetica", 10)

    for text in paragraphs:
        for line in _wrap_pdf_text(text, max_chars=105):
            if y < 20 * mm:
                pdf.showPage()
                pdf.setFont("Helvetica", 10)
                y = height - 20 * mm
            pdf.drawString(20 * mm, y, line)
            y -= 5 * mm
        y -= 3 * mm

    pdf.save()
    return out.getvalue()


def _wrap_pdf_text(text: str, max_chars: int) -> list[str]:
    words = text.split()
    if not words:
        return [""]
    lines: list[str] = []
    current = words[0]
    for word in words[1:]:
        candidate = f"{current} {word}"
        if len(candidate) <= max_chars:
            current = candidate
        else:
            lines.append(current)
            current = word
    lines.append(current)
    return lines


def generate_pptx(*, title: str, slides: list[dict[str, Any]]) -> bytes:
    title = _ensure_text(title, "title").strip()
    if not title:
        raise ValueError("title cannot be empty")
    if not isinstance(slides, list) or not slides or len(slides) > MAX_SLIDES:
        raise ValueError(f"slides must be a non-empty list with at most {MAX_SLIDES} slides")

    from pptx import Presentation

    prs = Presentation()
    first_layout = prs.slide_layouts[0]
    title_slide = prs.slides.add_slide(first_layout)
    title_slide.shapes.title.text = title

    text_budget = len(title.encode("utf-8"))
    bullet_layout = prs.slide_layouts[1]
    for spec in slides:
        if not isinstance(spec, dict):
            raise ValueError("Each slide must be an object")
        slide_title = _ensure_text(spec.get("title", ""), "slide title")
        bullets = spec.get("bullets", [])
        if not isinstance(bullets, list) or len(bullets) > MAX_BULLETS_PER_SLIDE:
            raise ValueError(f"bullets must be a list with at most {MAX_BULLETS_PER_SLIDE} items")
        slide = prs.slides.add_slide(bullet_layout)
        slide.shapes.title.text = slide_title
        body = slide.placeholders[1].text_frame
        body.clear()
        for index, bullet in enumerate(bullets):
            bullet_text = _ensure_text(bullet, "bullet")
            paragraph = body.paragraphs[0] if index == 0 else body.add_paragraph()
            paragraph.text = bullet_text
            paragraph.level = 0
            text_budget += len(slide_title.encode("utf-8")) + len(bullet_text.encode("utf-8"))
            if text_budget > MAX_TEXT_BYTES:
                raise ValueError(f"Generated text exceeds {MAX_TEXT_BYTES} bytes")

    out = BytesIO()
    prs.save(out)
    return out.getvalue()


def extension_for(kind: str) -> tuple[str, str]:
    mapping = {
        "docx": (".docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
        "xlsx": (".xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
        "pdf": (".pdf", "application/pdf"),
        "pptx": (".pptx", "application/vnd.openxmlformats-officedocument.presentationml.presentation"),
    }
    try:
        return mapping[kind]
    except KeyError as exc:
        raise ValueError(f"Unsupported generator kind: {kind}") from exc
